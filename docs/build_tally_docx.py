import openpyxl, re
from collections import Counter
X="/root/.claude/uploads/45f611ed-8eb7-593b-80cb-7e0fbd47b411/791b7c02-TrikRide_Survey_Tally.xlsx"
wb=openpyxl.load_workbook(X,data_only=True)
ITEM=re.compile(r'^[A-F]\d+$')
SECT={"Passenger Responses":[("A","Functional Suitability"),("B","Usability"),("C","Efficiency"),
                            ("D","Reliability"),("E","Security"),("F","Overall Satisfaction")],
      "Driver Responses":[("A","Usability"),("B","Functionality"),("C","Efficiency"),("D","Reliability")],
      "Admin Responses":[("A","Functional Suitability"),("B","Usability"),("C","Reliability")]}
TITLE={"Passenger Responses":("Student Passengers","N = 20"),
       "Driver Responses":("Tricycle Drivers","N = 20"),
       "Admin Responses":("System Administrator","N = 1")}
L=["# TrikRide — Survey Tally","",
   "Frequency of responses to each statement of the evaluation questionnaire.",
   "",
   "**Scale:** 5 = Strongly Agree · 4 = Agree · 3 = Neutral · 2 = Disagree · 1 = Strongly Disagree",
   "","[[PB]]"]
for sheet,(title,n) in TITLE.items():
    ws=wb[sheet]
    lab=list(next(ws.iter_rows(min_row=3,max_row=3,values_only=True)))
    cod=[str(c) if c else "" for c in next(ws.iter_rows(min_row=4,max_row=4,values_only=True))]
    rows=[r for r in ws.iter_rows(min_row=6,max_row=ws.max_row,values_only=True) if r[0] is not None]
    L+= [f"## {title} ({n})","",
         "| # | Statement | 5 | 4 | 3 | 2 | 1 | Total |",
         "|:---|:---|---:|---:|---:|---:|---:|---:|"]
    for letter,name in SECT[sheet]:
        L.append(f"| | **{letter}. {name}** | | | | | | |")
        k=0
        for i in range(len(cod)):
            if not (ITEM.match(cod[i]) and cod[i].startswith(letter) and lab[i]): continue
            k+=1
            v=[r[i] for r in rows if isinstance(r[i],int)]; c=Counter(v)
            stmt=str(lab[i]).strip().replace("|","\\|")
            L.append(f"| {k} | {stmt} | {c[5]} | {c[4]} | {c[3]} | {c[2]} | {c[1]} | {len(v)} |")
    L+= ["","[[PB]]"]
open("tally_only.md","w").write("\n".join(L)+"\n")
print("\n".join(L[:26]))
